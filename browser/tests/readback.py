"""Independent package/XML and openpyxl checks of the browser writer's outputs."""
import importlib.util
import io
import json
from pathlib import Path
import re
import zipfile
from xml.etree import ElementTree as ET
import openpyxl

ROOT=Path(__file__).resolve().parents[2]
spec=importlib.util.spec_from_file_location('handoff',ROOT/'scripts/csv_handoff.py')
cli=importlib.util.module_from_spec(spec);spec.loader.exec_module(cli)
NS={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
for name in ['fictional-sample','edge-values']:
    folder=ROOT/'browser/.fixtures'
    expected=json.loads((folder/f'{name}.json').read_text())
    data=(folder/f'{name}.xlsx').read_bytes()
    assert len(data)<=10*1024*1024
    actual=cli.read_xlsx(data)
    assert [[c['value'] for c in row] for row in actual]==expected
    assert all(c['type']=='text' for row in actual for c in row)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert len(z.infolist())==6
        assert sum(i.file_size for i in z.infolist())<=40*1024*1024
        assert all(i.file_size<=20*1024*1024 and i.file_size<=max(1,i.compress_size)*1000 for i in z.infolist())
        sheet=ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
        assert sheet.find('.//s:f',NS) is None
        cells=sheet.findall('.//s:sheetData/s:row/s:c',NS)
        assert len(cells)==sum(map(len,expected))
        styles=ET.fromstring(z.read('xl/styles.xml')).find('s:cellXfs',NS)
        decoded=[]
        for row in sheet.findall('s:sheetData/s:row',NS):
            values=[]
            for cell in row:
                assert cell.attrib['t']=='inlineStr'
                assert styles[int(cell.attrib['s'])].attrib['numFmtId']=='49'
                text=cell.find('s:is/s:t',NS).text or ''
                values.append(re.sub(r'_x([0-9a-fA-F]{4})_',lambda m:chr(int(m[1],16)),text))
            decoded.append(values)
        assert decoded==expected
    wb=openpyxl.load_workbook(io.BytesIO(data),data_only=False)
    assert wb.sheetnames==['Data'] and wb['Data'].freeze_panes=='A2'
    for r,row in enumerate(expected,1):
        for c,value in enumerate(row,1):
            cell=wb['Data'].cell(r,c)
            assert cell.data_type=='s' and cell.number_format=='@'
            assert cell.hyperlink is None
            # Existing known openpyxl behavior: escaped inline literals expose the prefix.
            # CLI + direct OOXML checks above assert those exact values instead.
            if not re.search(r'_x[0-9a-fA-F]{4}_',value):
                assert (cell.value or '')==value,(name,r,c,repr(cell.value),repr(value))
    print(f'{name}: exact CLI readback, independent XML checks, and openpyxl ordinary-value/type checks passed')
print('Not a Microsoft Excel or stable LibreOffice application test.')
